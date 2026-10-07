"""The apply kit: a keyword-tailored resume and every question answered, all grounded."""

import json
from datetime import UTC, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from agent import apply_kit, dashboard, resume_tailor
from agent.answer_agent import KitAnswer, answer_questions, questions_for_job
from agent.answers import AnswerDecision
from agent.library import LibraryDocument
from agent.resume_tailor import (
    KEYWORD_PROMPT,
    POOL_PROMPT,
    Entry,
    Keywords,
    TailoredResume,
    bullet_problem,
    extract_keywords,
    keyword_report,
    tailor_resume,
    write_docx,
)
from db.models import Application, Base, Job

RESUME = """Sam Sample
sam@example.com | github.com/sam
Example University, B.S. Computer Science, 2025
Research Assistant, Example Lab
Built a retrieval pipeline in Python that answered 120 questions a day."""
BANK = """# Projects
## Agent Evaluator
Wrote an evaluation harness in Python for LLM agents, scoring 40 tasks.
Used PyTorch to fine-tune a small ranking model."""
DESCRIPTION = (
    "<p>You will <b>evaluate</b> LLM agents and collaborate with researchers. "
    "Experience with Python, PyTorch, and Kubernetes. Build evaluation pipelines.</p>"
)
REAL_PAGE_COUNT = resume_tailor._page_count


@pytest.fixture(autouse=True)
def _estimated_layout(monkeypatch: pytest.MonkeyPatch) -> None:
    """Measure pages with the estimate, not LibreOffice, except where a test asks for it."""
    monkeypatch.setattr(resume_tailor, "_page_count", lambda _path: None)


JOB = {"company": "Acme", "title": "AI Engineer", "description": DESCRIPTION}


class FakeClient:
    """Answers each structured request from a script keyed by its system prompt."""

    def __init__(self, tailor_replies: list[dict]) -> None:
        self.tailor_replies = list(tailor_replies)
        self.requests: list[dict] = []
        self.messages = SimpleNamespace(create=self.create)
        self.beta = SimpleNamespace(messages=SimpleNamespace(create=self.create))

    def create(self, **request):
        self.requests.append(request)
        if request["system"] == KEYWORD_PROMPT:
            payload = {
                "action_verbs": ["Evaluate", "collaborate", "lead"],  # "lead" is not there
                "technologies": ["python", "PyTorch", "Kubernetes"],
                "concepts": ["evaluation pipelines"],
                "required": ["Python", "PyTorch"],
                "preferred": ["Kubernetes", "Rust"],  # Rust is not in the description
            }
        elif request["system"] == POOL_PROMPT:
            payload = {
                "name": "Sam Sample",
                "contact": ["sam@example.com", "github.com/sam"],
                "education": ["Example University, B.S. Computer Science, 2025"],
                "entries": [
                    {
                        "title": "Research Assistant", "organization": "Example Lab",
                        "location": "", "dates": "", "kind": "experience", "doc_id": "resume",
                        "facts": ["Built a retrieval pipeline in Python that answered 120 "
                                  "questions a day."],
                    },
                    {
                        "title": "Agent Evaluator", "organization": "", "location": "",
                        "dates": "", "kind": "project", "doc_id": "bank.md",
                        "facts": ["Wrote an evaluation harness in Python for LLM agents, "
                                  "scoring 40 tasks."],
                    },
                    {
                        "title": "Imaginary Startup", "organization": "", "location": "",
                        "dates": "", "kind": "experience", "doc_id": "resume", "facts": [],
                    },
                ],
            }
        elif request["system"] == resume_tailor.SUGGEST_PROMPT:
            payload = {"suggestions": []}
        else:
            payload = self.tailor_replies.pop(0)
        return SimpleNamespace(
            stop_reason="end_turn", content=[SimpleNamespace(type="text", text=json.dumps(payload))]
        )


def _bullet(text: str, *quotes: tuple[str, str]) -> dict:
    return {"text": text, "evidence": [{"doc_id": d, "quote": q} for d, q in quotes]}


GOOD = _bullet(
    "Evaluate LLM agents with an evaluation harness in Python, scoring 40 tasks",
    ("bank.md", "Wrote an evaluation harness in Python for LLM agents, scoring 40 tasks."),
)
CLAIMS_K8S = _bullet(
    "Deployed the retrieval pipeline on Kubernetes",
    ("resume", "Built a retrieval pipeline in Python"),
)


def test_keywords_are_kept_only_as_the_description_spells_them() -> None:
    keywords = extract_keywords(
        " ".join(DESCRIPTION.split()), client=FakeClient([]), model="m"
    )

    assert keywords.action_verbs == ["evaluate", "collaborate"]
    assert keywords.technologies == ["Python", "PyTorch", "Kubernetes"]
    assert keywords.concepts == ["evaluation pipelines"]
    assert keywords.required == ["Python", "PyTorch"]
    assert keywords.preferred == ["Kubernetes"]


def test_bullets_naming_unsupported_tools_or_numbers_are_rejected() -> None:
    docs = {"resume": LibraryDocument("resume", "Resume", RESUME)}
    keywords = Keywords(technologies=["Python", "Kubernetes"])

    assert "Kubernetes" in bullet_problem(CLAIMS_K8S, keywords, docs)
    inflated = _bullet(
        "Built a retrieval pipeline in Python serving 500 questions a day",
        ("resume", "Built a retrieval pipeline in Python that answered 120 questions a day."),
    )
    assert "500" in bullet_problem(inflated, keywords, docs)
    invented = _bullet("Built things", ("resume", "Built a rocket"))
    assert "not in candidate document" in bullet_problem(invented, keywords, docs)
    from_job = _bullet("Python", ("job_description", "Python"))
    assert bullet_problem(from_job, keywords, docs)


def test_the_tailor_asks_about_missing_keywords_and_keeps_only_sourced_bullets(
    tmp_path: Path,
) -> None:
    library = tmp_path / "library"
    library.mkdir()
    (library / "bank.md").write_text(BANK, encoding="utf-8")
    draft = {
        "sections": [
            {
                "heading": "Experience",
                "entries": [
                    {"entry_id": "e1", "bullets": [CLAIMS_K8S]},
                    {"entry_id": "e2", "bullets": [GOOD]},
                    {"entry_id": "e3", "bullets": [GOOD]},  # dropped from the pool
                ],
            }
        ],
        "skills": [{"label": "Tools", "items": ["Python", "PyTorch", "Rust"]}],
    }
    client = FakeClient([draft, draft])  # the fix round repeats the same mistake
    asked: list[str] = []

    resume = tailor_resume(
        JOB, RESUME, client=client, model="m", library_dir=library,
        ask=lambda keyword, _suggestion: asked.append(keyword) or None,
    )

    # Python and PyTorch are in the documents; the bank says "evaluation harness", not
    # "evaluation pipelines", so that is asked about too.
    assert asked == ["Kubernetes", "evaluation pipelines"]
    assert resume.gaps == ["Kubernetes", "evaluation pipelines"]
    sections = dict(resume.sections)
    assert list(sections) == ["Experience", "Technical Projects"]  # grouped by kind
    (lab,) = sections["Experience"]
    (evaluator,) = sections["Technical Projects"]
    assert lab.bullets == [
        "Built a retrieval pipeline in Python that answered 120 questions a day."
    ]  # the Kubernetes claim failed twice, so the original wording stays
    assert evaluator.bullets == [GOOD["text"]]
    assert dict(resume.skills) == {"Tools": ["Python", "PyTorch"]}  # no source mentions Rust
    assert "evaluate" in resume.used and "Kubernetes" not in resume.used
    fix_round = client.requests[-1]["messages"][-1]["content"]
    assert "Kubernetes" in fix_round
    out = write_docx(resume, tmp_path / "resume.docx")
    from docx import Document

    text = "\n".join(paragraph.text for paragraph in Document(str(out)).paragraphs)
    assert "Sam Sample" in text and GOOD["text"] in text


def test_an_answer_to_a_missing_keyword_becomes_evidence(tmp_path: Path) -> None:
    library = tmp_path / "library"
    library.mkdir()
    deployed = _bullet(
        "Deployed the evaluator on Kubernetes",
        ("keyword_answers.md", "Kubernetes: deployed the evaluator on a Kubernetes cluster"),
    )
    draft = {
        "sections": [
            {"heading": "Projects", "entries": [{"entry_id": "e2", "bullets": [deployed]}]}
        ],
        "skills": [],
    }
    (library / "bank.md").write_text(BANK, encoding="utf-8")

    resume = tailor_resume(
        JOB, RESUME, client=FakeClient([draft]), model="m", library_dir=library,
        ask=lambda keyword, _suggestion: (
            "deployed the evaluator on a Kubernetes cluster" if keyword == "Kubernetes" else None
        ),
    )

    assert resume.gaps == ["evaluation pipelines"]
    assert dict(resume.sections)["Technical Projects"][0].bullets == [deployed["text"]]
    saved = (library / "keyword_answers.md").read_text(encoding="utf-8")
    assert "Kubernetes: deployed the evaluator" in saved


def test_every_question_is_answered_including_the_common_ones() -> None:
    recorded = {
        "First name": "Sam",
        "Why are you interested in Acme?": None,
        "Agent summary": "Done.",
    }
    questions = questions_for_job("Acme", "AI Engineer", recorded)

    asked = [question for question, _ in questions]
    assert "Agent summary" not in asked
    assert "Why do you want to work at Acme?" not in asked  # the form already asks it
    assert "Why are you a good fit for the AI Engineer role?" in asked
    assert "Tell us about a technical project you are proud of." in asked

    def answerer(question: str, long_form: bool) -> AnswerDecision:
        if long_form:
            return AnswerDecision(f"Draft for {question}", "[resume] x", True, "Review.", True)
        return AnswerDecision(None, None, True, "No source.")

    answers = {item.question: item for item in answer_questions(questions, answerer)}
    assert answers["First name"] == KitAnswer("First name", "Sam", "answer",
                                              "Filled from your profile on the form.")
    assert answers["Why are you interested in Acme?"].kind == "draft"
    assert answers["Tell us about a technical project you are proud of."].kind == "draft"


@pytest.fixture
def session():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        yield db
    engine.dispose()


def test_the_kit_is_recorded_and_shown_in_the_apply_panel(
    session: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(apply_kit, "KITS_DIR", tmp_path / "kits")
    monkeypatch.setattr(dashboard, "KITS_DIR", tmp_path / "kits")
    job = Job(
        source="lever", platform="lever", company="Acme", title="AI Engineer",
        url="https://jobs.lever.co/acme/1", location_raw="Remote - US",
        location_category="remote_us", description=DESCRIPTION, status="new",
    )
    session.add(job)
    session.flush()
    session.add(
        Application(job_id=job.id, mode="dry_run", answers={"First name": "Sam", "Why Acme?": None})
    )
    session.flush()
    session.refresh(job)
    library = tmp_path / "library"
    library.mkdir()
    (library / "bank.md").write_text(BANK, encoding="utf-8")
    monkeypatch.setattr("agent.resume_tailor.LIBRARY_DIR", library)
    draft = {
        "sections": [{"heading": "Projects", "entries": [{"entry_id": "e2", "bullets": [GOOD]}]}],
        "skills": [],
    }

    kit = apply_kit.build_kit(
        session, job, profile={}, resume_text=RESUME, client=FakeClient([draft]),
        answer_model="m", resume_model="m", ask=None,
        answerer=lambda question, _long: AnswerDecision(f"Draft: {question}", "e", True, "r", True),
    )

    assert not kit.problems and kit.resume_path.is_file()
    session.refresh(job)
    stored = apply_kit.latest_kit(job)
    assert stored.answers["First name"] == "Sam"
    assert stored.suggested_answers["Why Acme?"] == "Draft: Why Acme?"
    assert "Kubernetes" in stored.field_notes["Resume keywords"]
    page = tmp_path / "dashboard.html"
    dashboard.write_dashboard(session, page, fit_threshold=70)
    html = page.read_text(encoding="utf-8")
    assert f"href='/kits/{job.id}/resume'" in html
    assert "Draft: Tell us about a technical project you are proud of." in html
    assert "class='apply-link' target='_blank'" in html


def _job_with_attempts(session: Session, *attempts: Application) -> Job:
    job = Job(
        source="greenhouse", platform="greenhouse", company="Acme", title="AI Engineer",
        url="https://boards.greenhouse.io/acme/jobs/1", location_raw="Remote - US",
        location_category="remote_us", description=DESCRIPTION, status="manual_review",
    )
    session.add(job)
    session.flush()
    start = datetime(2026, 10, 1, tzinfo=UTC)
    for index, attempt in enumerate(attempts):
        attempt.job_id = job.id
        attempt.started_at = start + timedelta(hours=index)
        session.add(attempt)
    session.flush()
    session.refresh(job)
    return job


def test_questions_from_every_form_reading_are_kept(session: Session) -> None:
    job = _job_with_attempts(
        session,
        Application(mode="dry_run", answers={"First name": "Sam", "Why Acme?": None}),
        Application(mode="hand_off", answers={"First name": "Sam", "Visa status?": None}),
    )

    recorded = apply_kit._form_answers(job)

    assert recorded == {"First name": "Sam", "Why Acme?": None, "Visa status?": None}


def test_questions_found_after_the_kit_still_show_in_the_apply_panel(
    session: Session, tmp_path: Path
) -> None:
    job = _job_with_attempts(
        session,
        Application(mode="dry_run", answers={"First name": "Sam"}),
        Application(
            mode=apply_kit.KIT_MODE,
            answers={"First name": "Sam", "Why do you want to work at Acme?": None},
            suggested_answers={"Why do you want to work at Acme?": "Draft: why Acme"},
        ),
        Application(
            mode="hand_off",
            answers={"First name": "Sam", "Preferred pronouns": "they/them",
                     "Describe a hard bug you fixed.": None},
        ),
    )
    kit = apply_kit.latest_kit(job)

    assert apply_kit.questions_missing_from_kit(job, kit) == [
        ("Preferred pronouns", "they/them"),
        ("Describe a hard bug you fixed.", None),
    ]
    page = tmp_path / "dashboard.html"
    dashboard.write_dashboard(session, page, fit_threshold=70)
    html = page.read_text(encoding="utf-8")
    assert "Draft: why Acme" in html
    assert "Describe a hard bug you fixed." in html
    assert "they/them" in html
    assert dashboard.MISSING_FROM_KIT_NOTE in html
    assert dashboard.FORM_NOT_READ_NOTE not in html


def test_the_apply_panel_says_when_the_form_was_never_read(
    session: Session, tmp_path: Path
) -> None:
    _job_with_attempts(
        session,
        Application(
            mode=apply_kit.KIT_MODE,
            answers={"Why do you want to work at Acme?": None},
            suggested_answers={"Why do you want to work at Acme?": "Draft: why Acme"},
        ),
    )

    page = tmp_path / "dashboard.html"
    dashboard.write_dashboard(session, page, fit_threshold=70)

    assert dashboard.FORM_NOT_READ_NOTE in page.read_text(encoding="utf-8")


def test_a_kit_without_a_form_reading_answers_the_boards_questions(
    session: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(apply_kit, "KITS_DIR", tmp_path / "kits")
    monkeypatch.setattr(dashboard, "KITS_DIR", tmp_path / "kits")
    job = _job_with_attempts(session)
    monkeypatch.setattr(
        apply_kit, "board_questions", lambda _job: ["Why Example Robotics?", "Pronouns"]
    )

    apply_kit.build_kit(
        session, job, profile={}, resume_text=RESUME, client=None, answer_model="m",
        resume_model="m", make_resume=False,
        answerer=lambda question, _long: AnswerDecision(f"Answer: {question}", "e", False),
    )

    session.refresh(job)
    stored = apply_kit.latest_kit(job)
    assert stored.answers["Why Example Robotics?"] == "Answer: Why Example Robotics?"
    assert stored.answers["Pronouns"] == "Answer: Pronouns"
    assert stored.field_notes[apply_kit.BOARD_QUESTIONS_NOTE]
    page = tmp_path / "dashboard.html"
    dashboard.write_dashboard(session, page, fit_threshold=70)
    html = page.read_text(encoding="utf-8")
    assert "Answer: Pronouns" in html
    assert dashboard.FORM_NOT_READ_NOTE not in html


def test_board_questions_are_read_only_from_greenhouse_links(
    session: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    calls = []
    monkeypatch.setattr(
        "agent.fetchers.greenhouse.fetch_greenhouse_questions",
        lambda board, job_id: calls.append((board, job_id)) or ["Why Acme?"],
    )
    job = _job_with_attempts(session)

    assert apply_kit.board_questions(job) == ["Why Acme?"]
    assert calls == [("acme", "1")]
    job.platform = "lever"
    assert apply_kit.board_questions(job) == []


def test_a_failed_board_request_leaves_the_common_questions(
    session: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    def fail(_board: str, _job_id: str) -> list[str]:
        raise httpx.ConnectError("offline")

    monkeypatch.setattr("agent.fetchers.greenhouse.fetch_greenhouse_questions", fail)

    assert apply_kit.board_questions(_job_with_attempts(session)) == []


def _long_resume(entries: int = 6, bullets: int = 4) -> TailoredResume:
    filler = "and wrote it up for the team so others could reuse the approach later on"
    sections = [
        (
            heading,
            [
                Entry(
                    f"{heading[0]}{index}", f"{heading} {index}", "Example Org", "Remote",
                    "2024 - 2025", "project" if heading == "Technical Projects" else "experience",
                    bullets=[f"Built tool {index}.{n} in Go {filler}" for n in range(bullets)],
                )
                for index in range(entries)
            ],
        )
        for heading in ("Experience", "Technical Projects")
    ]
    # One bullet uses a required qualification: it must survive the trimming.
    sections[1][1][-1].bullets[-1] = f"Evaluated LLM agents in Python {filler}"
    return TailoredResume(
        "Jordan Example", ["jordan@example.com", "Raleigh, NC"],
        ["Example State University, B.S. Computer Science, 2025"], sections,
        [("Languages", ["Python", "Go"])],
        Keywords(technologies=["Python", "Go"], required=["Python"]),
    )


def test_a_long_resume_is_trimmed_to_one_page_keeping_required_qualifications(
    tmp_path: Path,
) -> None:
    resume = _long_resume()

    write_docx(resume, tmp_path / "resume.docx")

    assert resume_tailor._fits(resume, resume_tailor.MIN_FONT_SIZE)
    kept = [bullet for _title, entries in resume.sections for e in entries for bullet in e.bullets]
    assert any("Evaluated LLM agents in Python" in bullet for bullet in kept)
    assert len(kept) < 48
    assert resume.notes[-1].startswith("To fit one page:")


def test_the_resume_uses_times_new_roman_at_10pt_or_more_in_four_sections(
    tmp_path: Path,
) -> None:
    from docx import Document

    short = _long_resume(entries=1, bullets=2)

    out = write_docx(short, tmp_path / "resume.docx")

    document = Document(str(out))
    runs = [run for paragraph in document.paragraphs for run in paragraph.runs if run.text]
    assert {run.font.name for run in runs} == {"Times New Roman"}
    assert min(run.font.size.pt for run in runs) >= 10
    assert max(run.font.size.pt for run in runs if not run.bold) == 11  # roomy: largest size
    headings = [p.text for p in document.paragraphs if p.text.isupper()]
    assert headings == ["EDUCATION", "EXPERIENCE", "TECHNICAL PROJECTS", "SKILLS"]
    assert not short.notes  # nothing was trimmed


@pytest.mark.skipif(
    not __import__("shutil").which("soffice"),
    reason="LibreOffice is not installed",
)
def test_the_word_file_really_is_one_page(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(resume_tailor, "_page_count", REAL_PAGE_COUNT)
    resume = _long_resume()

    out = write_docx(resume, tmp_path / "resume.docx")

    assert REAL_PAGE_COUNT(out) == 1
    assert any("Evaluated LLM agents in Python" in b for _t, es in resume.sections
               for e in es for b in e.bullets)


def test_the_keyword_report_counts_required_and_preferred_matches() -> None:
    resume = _long_resume(entries=1, bullets=1)
    resume.keywords = Keywords(required=["Python", "Rust"], preferred=["Go"])

    report = keyword_report(resume)

    assert "Required qualifications matched: 1 of 2 (missing: Rust)" in report
    assert "Preferred qualifications matched: 1 of 1" in report
